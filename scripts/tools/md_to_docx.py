#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Convert project Markdown docs to Word (.docx)."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_doc_default_font(doc: Document, name: str = "\u5fae\u8f6f\u96c5\u9ed1", size: int = 11) -> None:
    style = doc.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(size)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_inline_runs(paragraph, text: str, bold_default: bool = False) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos : m.start()])
            run.bold = bold_default
        chunk = m.group(0)
        if chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        elif chunk.startswith("["):
            link_m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", chunk)
            if link_m:
                label, url = link_m.groups()
                run = paragraph.add_run(f"{label} ({url})")
                run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
                run.underline = True
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.bold = bold_default


def is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and "|" in s[1:-1]


def is_table_sep(line: str) -> bool:
    s = line.strip().strip("|")
    if not s:
        return False
    parts = [p.strip() for p in s.split("|")]
    return all(re.fullmatch(r":?-{3,}:?", p) for p in parts if p)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and is_table_row(lines[i]):
        if is_table_sep(lines[i]):
            i += 1
            continue
        row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        rows.append(row)
        i += 1
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = table.rows[ri].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            text = row[ci] if ci < len(row) else ""
            if ri == 0:
                add_inline_runs(p, text, bold_default=True)
                set_cell_shading(cell, "E7E6E6")
            else:
                add_inline_runs(p, text)
    doc.add_paragraph()


def convert_md_to_docx(md_path: Path, docx_path: Path) -> None:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    set_doc_default_font(doc)

    title = md_path.stem
    for line in lines:
        if line.startswith("# "):
            title = line[2:].strip()
            break

    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run(title)
    tr.bold = True
    tr.font.size = Pt(18)
    doc.add_paragraph()

    i = 0
    in_code = False
    code_buf: list[str] = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.style = "No Spacing"
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                run.font.size = Pt(9)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph("\u2500" * 40)
            i += 1
            continue

        if is_table_row(stripped):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
            i += 1
            continue
        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=4)
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            add_inline_runs(p, stripped[2:])
            i += 1
            continue

        num_m = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if num_m:
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, num_m.group(2))
            i += 1
            continue

        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, stripped[2:])
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, stripped)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def main() -> int:
    root = Path(__file__).resolve().parents[2]
    docs_dir = root / "docs"
    out_dir = docs_dir / "word"

    files = [
        ("01-\u9700\u6c42\u8bf4\u660e\u4e66.md", "01-\u9700\u6c42\u8bf4\u660e\u4e66.docx"),
        ("02-\u6982\u8981\u8bbe\u8ba1.md", "02-\u6982\u8981\u8bbe\u8ba1.docx"),
        ("03-\u7cfb\u7edf\u8bbe\u8ba1.md", "03-\u8be6\u7ec6\u8bbe\u8ba1.docx"),
        ("05-\u5b9e\u65bd\u5907\u5fd8.md", "05-\u5b9e\u65bd\u5907\u5fd8.docx"),
        ("06-\u4f7f\u7528\u8bf4\u660e\u4e66.md", "06-\u4f7f\u7528\u8bf4\u660e\u4e66.docx"),
    ]

    for src_name, dst_name in files:
        src = docs_dir / src_name
        dst = out_dir / dst_name
        if not src.exists():
            print(f"SKIP missing: {src_name}", file=sys.stderr)
            continue
        convert_md_to_docx(src, dst)
        print(f"OK: {dst}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
